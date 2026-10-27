
import re
from pathlib import Path
from typing import Iterable

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import PyPDFLoader, TextLoader
from docx import Document as DocxDocument


SUPPORTED = {".pdf", ".txt", ".md", ".docx"}


def load_file(path: Path) -> list[Document]:
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return PyPDFLoader(str(path)).load()

    if suffix in {".txt", ".md"}:
        return TextLoader(str(path), encoding="utf-8").load()

    if suffix == ".docx":
        doc = DocxDocument(str(path))
        text = "\n".join(
            p.text for p in doc.paragraphs if p.text.strip()
        )
        return [
            Document(
                page_content=text,
                metadata={"source": str(path)},
            )
        ]

    raise ValueError(f"Unsupported file type: {suffix}")


def split_markdown_sections(doc: Document) -> list[Document]:
    """
    Split Markdown documents by headings and retain section metadata.
    Non-Markdown documents pass through unchanged.
    """
    source = doc.metadata.get("source", "")
    if Path(source).suffix.lower() != ".md":
        return [doc]

    lines = doc.page_content.splitlines()
    sections = []
    heading_stack = []
    current_lines = []

    for line in lines:
        match = re.match(r"^(#{1,6})\s+(.+?)\s*#*\s*$", line)

        if match:
            if current_lines:
                content = "\n".join(current_lines).strip()
                if content:
                    sections.append((heading_stack.copy(), content))
                current_lines = []

            level = len(match.group(1))
            heading = match.group(2).strip()

            heading_stack = heading_stack[:level - 1]
            heading_stack.append(heading)

            current_lines.append(line)
        else:
            current_lines.append(line)

    if current_lines:
        content = "\n".join(current_lines).strip()
        if content:
            sections.append((heading_stack.copy(), content))

    return [
        Document(
            page_content=content,
            metadata={
                **doc.metadata,
                "section": headings[-1] if headings else "Unsectioned",
                "section_path": " > ".join(headings),
            },
        )
        for headings, content in sections
    ]


def chunk_documents(docs: Iterable[Document]) -> list[Document]:
    sectioned_docs = []

    for doc in docs:
        sectioned_docs.extend(split_markdown_sections(doc))

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=900,
        chunk_overlap=120,
        add_start_index=True,
    )

    return splitter.split_documents(sectioned_docs)
